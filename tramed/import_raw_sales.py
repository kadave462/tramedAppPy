"""
Run once to load the Ishyiga Excel export into the raw_sales table.

Usage:
    python import_raw_sales.py "path/to/sales of 13_05.xlsx"

If no path is given, defaults to the educational materials folder.
"""

import sys
import os
from openpyxl import load_workbook

# ── bootstrap Flask app context so SQLAlchemy models work ──────────────────
from app import app
from database import db, RawSale

DEFAULT_EXCEL = os.path.join(
    os.path.dirname(__file__),
    "..",
    "educational materials",
    "Phase 1 plan",
    "sales of 13_05.xlsx",
)

EXCEL_PATH = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_EXCEL


def cell_val(row, idx, cast=None):
    """Return the value of the cell at 1-based index, optionally cast."""
    val = row[idx - 1].value
    if val is None:
        return None
    if cast:
        try:
            return cast(val)
        except (ValueError, TypeError):
            return None
    return val


def import_excel(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(min_row=2))   # skip header row
    if not rows:
        print("No data rows found.")
        return

    inserted = 0
    skipped  = 0

    with app.app_context():
        db.create_all()   # creates raw_sales if it doesn't exist yet

        for row in rows:
            # Skip completely empty rows
            if all(c.value is None for c in row):
                skipped += 1
                continue

            raw = RawSale(
                id_invoice       = cell_val(row, 1,  int),
                date             = str(cell_val(row, 2))  if cell_val(row, 2)  is not None else None,
                code             = str(cell_val(row, 3))  if cell_val(row, 3)  is not None else None,
                name_product     = str(cell_val(row, 4))  if cell_val(row, 4)  is not None else None,
                num_lot          = str(cell_val(row, 5))  if cell_val(row, 5)  is not None else None,
                quantite         = cell_val(row, 6,  int),
                price            = cell_val(row, 7,  float),
                price_revient    = cell_val(row, 8,  float),
                total            = cell_val(row, 9,  float),
                tva              = cell_val(row, 10, float),
                employe          = str(cell_val(row, 11)) if cell_val(row, 11) is not None else None,
                num_affiliation  = str(cell_val(row, 12)) if cell_val(row, 12) is not None else None,
                numero_quittance = str(cell_val(row, 13)) if cell_val(row, 13) is not None else None,
                nom_client       = str(cell_val(row, 14)) if cell_val(row, 14) is not None else None,
                prenom_client    = str(cell_val(row, 15)) if cell_val(row, 15) is not None else None,
                percentage       = cell_val(row, 16, float),
                time             = str(cell_val(row, 17)) if cell_val(row, 17) is not None else None,
                type_paiement    = str(cell_val(row, 18)) if cell_val(row, 18) is not None else None,
                monnaie          = cell_val(row, 19, float),
            )
            db.session.add(raw)
            inserted += 1

        db.session.commit()

    print(f"Done — {inserted} rows inserted, {skipped} empty rows skipped.")


if __name__ == "__main__":
    import_excel(EXCEL_PATH)
